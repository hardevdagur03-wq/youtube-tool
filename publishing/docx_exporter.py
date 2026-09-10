from __future__ import annotations

import hashlib
import logging
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls

from publishing.models import (
    DocumentModel, ExportFileInfo, ExportFormat, FORMAT_EXTENSIONS, FORMAT_MIME_TYPES,
)

logger = logging.getLogger(__name__)

BRAND_COLORS = {
    "primary": RGBColor(0x05, 0x96, 0x69),
    "secondary": RGBColor(0x1F, 0x29, 0x37),
    "accent": RGBColor(0x7C, 0x3A, 0xED),
    "text": RGBColor(0x1A, 0x1A, 0x2E),
    "muted": RGBColor(0x6B, 0x72, 0x80),
    "border": RGBColor(0xE5, 0xE7, 0xEB),
    "code_bg": RGBColor(0xF3, 0xF4, 0xF6),
}


class DocxExporter:
    def export(self, document: DocumentModel, output_dir: Path, base_url: str = "") -> ExportFileInfo:
        doc = Document()

        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)
        style.font.color.rgb = BRAND_COLORS["text"]
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15

        for level in range(1, 7):
            heading_style = doc.styles[f'Heading {level}']
            heading_style.font.name = 'Calibri'
            heading_style.font.color.rgb = BRAND_COLORS["secondary"]
            if level == 1:
                heading_style.font.size = Pt(22)
            elif level == 2:
                heading_style.font.size = Pt(16)
            elif level == 3:
                heading_style.font.size = Pt(13)
            else:
                heading_style.font.size = Pt(11)

        section = doc.sections[0]
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

        self._add_cover_page(doc, document)
        doc.add_page_break()

        self._add_toc(doc, document)
        doc.add_page_break()

        self._add_metadata_block(doc, document)

        self._add_introduction(doc, document)

        for section_info in document.sections:
            self._add_section(doc, section_info)

        if document.faq:
            self._add_faq(doc, document)

        if document.conclusion:
            self._add_conclusion(doc, document)

        if document.call_to_action:
            self._add_cta(doc, document)

        if document.references:
            self._add_references(doc, document)

        footer = section.footer
        footer.is_linked_to_previous = False
        footer_para = footer.paragraphs[0]
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer_para.add_run()
        fldChar1 = run._r.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'begin'})
        run._r.append(fldChar1)
        run2 = footer_para.add_run()
        instrText = run2._r.makeelement(qn('w:instrText'), {})
        instrText.text = ' PAGE '
        run2._r.append(instrText)
        run3 = footer_para.add_run()
        fldChar2 = run3._r.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'end'})
        run3._r.append(fldChar2)

        filename = f"blog{FORMAT_EXTENSIONS[ExportFormat.DOCX]}"
        filepath = output_dir / filename
        doc.save(str(filepath))
        size = filepath.stat().st_size
        checksum = hashlib.sha256(filepath.read_bytes()).hexdigest()[:32]
        return ExportFileInfo(
            format=ExportFormat.DOCX,
            filename=filename,
            size_bytes=size,
            size_display=self._size_display(size),
            mime_type=FORMAT_MIME_TYPES[ExportFormat.DOCX],
            checksum=checksum,
        )

    def _add_cover_page(self, doc: Document, document: DocumentModel) -> None:
        for _ in range(6):
            doc.add_paragraph("")

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(document.seo_title or document.title)
        run.font.size = Pt(28)
        run.font.color.rgb = BRAND_COLORS["primary"]
        run.bold = True

        doc.add_paragraph("")

        if document.meta_description:
            desc = doc.add_paragraph()
            desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = desc.add_run(document.meta_description)
            run.font.size = Pt(12)
            run.font.color.rgb = BRAND_COLORS["muted"]
            run.italic = True

        doc.add_paragraph("")

        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_text = ""
        if document.author:
            meta_text += f"By {document.author}"
        if document.publish_date:
            meta_text += f"  |  {document.publish_date}"
        if document.reading_time_minutes:
            meta_text += f"  |  {document.reading_time_minutes} min read"
        if meta_text:
            run = meta.add_run(meta_text)
            run.font.size = Pt(10)
            run.font.color.rgb = BRAND_COLORS["muted"]

    def _add_toc(self, doc: Document, document: DocumentModel) -> None:
        doc.add_heading("Table of Contents", level=1)
        for heading in document.headings:
            indent = "    " * (heading.level - 1) if heading.level > 1 else ""
            p = doc.add_paragraph(f"{indent}{heading.text}")
            p.paragraph_format.space_after = Pt(2)
            if heading.level == 1:
                p.runs[0].bold = True

    def _add_metadata_block(self, doc: Document, document: DocumentModel) -> None:
        rows = []
        if document.category:
            rows.append(("Category", document.category))
        if document.tags:
            rows.append(("Tags", ", ".join(document.tags)))
        if document.primary_keyword:
            rows.append(("Primary Keyword", document.primary_keyword))
        if document.word_count:
            rows.append(("Word Count", str(document.word_count)))
        if rows:
            table = doc.add_table(rows=len(rows), cols=2)
            table.style = 'Light Grid Accent 1'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, (key, val) in enumerate(rows):
                table.rows[i].cells[0].text = key
                table.rows[i].cells[1].text = val
                for cell in table.rows[i].cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after = Pt(2)
            doc.add_paragraph("")

    def _add_introduction(self, doc: Document, document: DocumentModel) -> None:
        if document.introduction:
            doc.add_heading("Introduction", level=1)
            for block in document.introduction.split("\n\n"):
                if block.strip():
                    self._add_text_block(doc, block.strip())

    def _add_section(self, doc: Document, section: Any) -> None:
        tag = section.heading_tag or "h2"
        level = int(tag[1]) if tag.startswith("h") else 2
        doc.add_heading(section.heading, level=level)

        if section.content:
            for block in section.content.split("\n\n"):
                if block.strip():
                    self._add_text_block(doc, block.strip())

        for sub in section.subsections:
            sub_level = min(level + 1, 6)
            doc.add_heading(sub.heading, level=sub_level)
            if sub.content:
                for block in sub.content.split("\n\n"):
                    if block.strip():
                        self._add_text_block(doc, block.strip())

    def _add_faq(self, doc: Document, document: DocumentModel) -> None:
        doc.add_heading("Frequently Asked Questions", level=1)
        for faq in document.faq:
            p = doc.add_paragraph()
            run = p.add_run(faq.question)
            run.bold = True
            run.font.color.rgb = BRAND_COLORS["secondary"]
            for block in faq.answer.split("\n\n"):
                if block.strip():
                    self._add_text_block(doc, block.strip())

    def _add_conclusion(self, doc: Document, document: DocumentModel) -> None:
        doc.add_heading("Conclusion", level=1)
        for block in document.conclusion.split("\n\n"):
            if block.strip():
                self._add_text_block(doc, block.strip())

    def _add_cta(self, doc: Document, document: DocumentModel) -> None:
        doc.add_heading("Call to Action", level=1)
        p = doc.add_paragraph()
        run = p.add_run(document.call_to_action)
        run.font.color.rgb = BRAND_COLORS["primary"]
        run.bold = True

    def _add_references(self, doc: Document, document: DocumentModel) -> None:
        doc.add_heading("References", level=1)
        for i, ref in enumerate(document.references, 1):
            p = doc.add_paragraph(f"{i}. {ref.label}")
            run = p.add_run(f" - {ref.url}")
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            run.underline = True

    def _add_text_block(self, doc: Document, text: str) -> None:
        lines = text.split("\n")
        for line in lines:
            line = line.strip()
            if not line:
                continue

            if line.startswith("- ") or line.startswith("* "):
                p = doc.add_paragraph(line[2:], style='List Bullet')
                continue

            if re.match(r"^\d+\.\s", line):
                p = doc.add_paragraph(re.sub(r"^\d+\.\s", "", line), style='List Number')
                continue

            if line.startswith("```"):
                continue

            if line.startswith("|"):
                continue

            p = doc.add_paragraph()
            self._add_formatted_text(p, line)

    def _add_formatted_text(self, paragraph: Any, text: str) -> None:
        parts = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("*") and part.endswith("*"):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            elif part.startswith("`") and part.endswith("`"):
                run = paragraph.add_run(part[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
            elif part.startswith("[") and "]" in part and "(" in part:
                link_match = re.match(r"\[(.+?)\]\((.+?)\)", part)
                if link_match:
                    run = paragraph.add_run(link_match.group(1))
                    run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
                    run.underline = True
            else:
                paragraph.add_run(part)

    @staticmethod
    def _size_display(size_bytes: int) -> str:
        if size_bytes < 1024:
            return f"{size_bytes} B"
        elif size_bytes < 1024 * 1024:
            return f"{size_bytes / 1024:.1f} KB"
        else:
            return f"{size_bytes / (1024 * 1024):.1f} MB"
